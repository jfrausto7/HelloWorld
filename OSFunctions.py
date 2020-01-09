"""This is a header block comment
This script has a bunch of OS functions and stuff i wanted to try out
Lol
"""

"""Below is an example of document manipulation.
It uses the docx module
Remember to install the 'python-docx' one, not the other one.
This was a pain to get to work.
"""

# import the docx module first
import docx
import openpyxl
from openpyxl.styles import Font

# now assign a variable for it
doc = docx.Document("./testdoc.docx")

for para in doc.paragraphs:
    print(para.text)

#a new println just for aesthetics
print("\n")

#openpyxl stuff below

wb = openpyxl.load_workbook("testbook.xlsx")

sheet = wb["Epic"]

#now for a for statement that iterates through the names in wb
for name in list(sheet.columns)[0]:
    if name.value == None:
        pass
    else:
        print(name.value)

#now for a font object and a change for cell A1
fontObject = Font("Avenir", size = 18, bold = True)
sheet["A1"].font = fontObject

#formula time
sheet["D5"] = "SUM"
fontObject2 = Font("Times New Roman", size = 18, underline = "single")
sheet["D5"].font = fontObject2
sheet["D6"] = "=SUM(B1:B5)"
#save the workbook
wb.save("testbook.xlsx")